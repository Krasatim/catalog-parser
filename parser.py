"""Парсер каталога товаров с выгрузкой в Excel, CSV или Google Таблицы.

Демо-источник: https://books.toscrape.com — учебный магазин, созданный
специально для практики парсинга.

Примеры:
    python parser.py                         # весь каталог -> books.xlsx
    python parser.py --pages 3 --out books.csv
    python parser.py --details               # + категория, остаток, описание
    python parser.py --details --sheet <ссылка на Google Таблицу>
"""

from __future__ import annotations

import argparse
import csv
import logging
import re
import time
from concurrent.futures import ThreadPoolExecutor
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

BASE_URL = "https://books.toscrape.com/catalogue/"
START_URL = urljoin(BASE_URL, "page-1.html")
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

log = logging.getLogger("parser")


@dataclass
class Product:
    title: str
    price: float
    currency: str
    rating: int
    in_stock: bool
    url: str
    image_url: str
    # Заполняются только с флагом --details
    category: str = ""
    upc: str = ""
    stock_qty: int | None = None
    description: str = ""


# Заголовки колонок в итоговой таблице
COLUMNS = {
    "title": "Название",
    "price": "Цена",
    "currency": "Валюта",
    "rating": "Рейтинг (1-5)",
    "in_stock": "В наличии",
    "url": "Ссылка",
    "image_url": "Изображение",
    "category": "Категория",
    "upc": "Артикул (UPC)",
    "stock_qty": "Остаток, шт.",
    "description": "Описание",
}


def make_session() -> requests.Session:
    """Сессия с повторными попытками при сетевых ошибках и 5xx/429."""
    session = requests.Session()
    retry = Retry(total=5, backoff_factor=1, status_forcelist=[429, 500, 502, 503, 504])
    session.mount("https://", HTTPAdapter(max_retries=retry, pool_maxsize=16))
    session.headers["User-Agent"] = (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/130.0 Safari/537.36"
    )
    return session


def fetch(session: requests.Session, url: str) -> BeautifulSoup:
    response = session.get(url, timeout=20)
    response.raise_for_status()
    # Передаём байты: lxml сам возьмёт кодировку из <meta charset>
    return BeautifulSoup(response.content, "lxml")


def parse_price(text: str) -> tuple[float, str]:
    match = re.search(r"([^\d\s.,]*)\s*([\d.,]+)", text)
    if not match:
        return 0.0, ""
    return float(match.group(2).replace(",", "")), match.group(1)


def parse_listing(soup: BeautifulSoup, page_url: str) -> list[Product]:
    products = []
    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        price, currency = parse_price(card.select_one(".price_color").get_text(strip=True))
        rating_class = card.select_one("p.star-rating")["class"]
        products.append(
            Product(
                title=link["title"],
                price=price,
                currency=currency,
                rating=next((RATINGS[c] for c in rating_class if c in RATINGS), 0),
                in_stock="In stock" in card.select_one(".availability").get_text(),
                url=urljoin(page_url, link["href"]),
                image_url=urljoin(page_url, card.select_one("img")["src"]),
            )
        )
    return products


def scrape_listing(session: requests.Session, max_pages: int | None, delay: float) -> list[Product]:
    """Обходит страницы каталога по ссылке «next», пока они не закончатся."""
    url: str | None = START_URL
    products: list[Product] = []
    page = 0
    while url and (max_pages is None or page < max_pages):
        page += 1
        soup = fetch(session, url)
        products.extend(parse_listing(soup, url))
        log.info("Страница %d: всего товаров %d", page, len(products))
        next_link = soup.select_one("li.next a")
        url = urljoin(url, next_link["href"]) if next_link else None
        if url:
            time.sleep(delay)
    return products


def add_details(session: requests.Session, product: Product) -> Product:
    """Дополняет товар данными с его карточки."""
    try:
        soup = fetch(session, product.url)
    except requests.RequestException as exc:
        log.warning("Не удалось загрузить %s: %s", product.url, exc)
        return product

    breadcrumbs = soup.select("ul.breadcrumb li")
    if len(breadcrumbs) >= 3:
        product.category = breadcrumbs[2].get_text(strip=True)

    table = {
        row.th.get_text(strip=True): row.td.get_text(strip=True)
        for row in soup.select("table.table tr")
    }
    product.upc = table.get("UPC", "")
    qty = re.search(r"\((\d+) available\)", table.get("Availability", ""))
    product.stock_qty = int(qty.group(1)) if qty else 0

    description = soup.select_one("#product_description + p")
    product.description = description.get_text(strip=True) if description else ""
    return product


def scrape_details(session: requests.Session, products: list[Product], workers: int) -> None:
    log.info("Загружаю карточки товаров (%d шт., потоков: %d)...", len(products), workers)
    with ThreadPoolExecutor(max_workers=workers) as pool:
        for done, _ in enumerate(pool.map(lambda p: add_details(session, p), products), 1):
            if done % 100 == 0 or done == len(products):
                log.info("Карточки: %d/%d", done, len(products))


def to_rows(products: list[Product], with_details: bool) -> tuple[list[str], list[list]]:
    keys = [f.name for f in fields(Product)]
    if not with_details:
        keys = keys[: keys.index("category")]
    header = [COLUMNS[k] for k in keys]
    rows = []
    for product in products:
        data = asdict(product)
        data["in_stock"] = "да" if data["in_stock"] else "нет"
        rows.append([data[k] for k in keys])
    return header, rows


def save_csv(path: Path, header: list[str], rows: list[list]) -> None:
    # utf-8-sig и «;» — чтобы файл корректно открывался в русском Excel
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(header)
        writer.writerows(rows)


def save_xlsx(path: Path, header: list[str], rows: list[list]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    ws.append(header)
    for row in rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {"Название": 45, "Ссылка": 30, "Изображение": 30, "Описание": 60}
    for idx, name in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 14)

    # Кликабельные ссылки
    for name in ("Ссылка", "Изображение"):
        if name in header:
            col = header.index(name) + 1
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=col)
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"

    wb.save(path)


def save_google_sheet(sheet: str, creds_path: Path, header: list[str], rows: list[list]) -> str:
    import gspread

    client = gspread.service_account(filename=str(creds_path))
    spreadsheet = client.open_by_url(sheet) if sheet.startswith("http") else client.open_by_key(sheet)
    ws = spreadsheet.sheet1
    ws.clear()
    ws.update([header] + [["" if v is None else v for v in row] for row in rows], value_input_option="USER_ENTERED")
    ws.format("1:1", {"textFormat": {"bold": True}})
    ws.freeze(rows=1)
    return spreadsheet.url


def main() -> None:
    parser = argparse.ArgumentParser(description="Парсер каталога товаров")
    parser.add_argument("--pages", type=int, help="сколько страниц каталога обойти (по умолчанию все)")
    parser.add_argument("--details", action="store_true", help="заходить в карточку каждого товара")
    parser.add_argument("--workers", type=int, default=8, help="потоков для загрузки карточек")
    parser.add_argument("--delay", type=float, default=0.3, help="пауза между страницами каталога, сек")
    parser.add_argument("--out", default="books.xlsx", help="файл результата: .xlsx или .csv")
    parser.add_argument("--sheet", help="ссылка или ID Google Таблицы для выгрузки")
    parser.add_argument("--creds", default="credentials.json", help="ключ сервисного аккаунта Google")
    args = parser.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s  %(message)s", datefmt="%H:%M:%S")
    started = time.perf_counter()

    session = make_session()
    products = scrape_listing(session, args.pages, args.delay)
    if args.details:
        scrape_details(session, products, args.workers)

    header, rows = to_rows(products, args.details)

    out = Path(args.out)
    if out.suffix.lower() == ".csv":
        save_csv(out, header, rows)
    else:
        save_xlsx(out, header, rows)
    log.info("Сохранено %d товаров в %s", len(rows), out.resolve())

    if args.sheet:
        url = save_google_sheet(args.sheet, Path(args.creds), header, rows)
        log.info("Выгружено в Google Таблицу: %s", url)

    log.info("Готово за %.1f сек", time.perf_counter() - started)


if __name__ == "__main__":
    main()
